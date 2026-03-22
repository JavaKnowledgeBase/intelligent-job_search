from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

BRAND_BLUE = RGBColor(0x2A, 0x93, 0xD5)
BRAND_GRAY = RGBColor(0x5A, 0x5A, 0x5D)
PDF_BRAND_BLUE = colors.HexColor("#2A93D5")
PDF_BRAND_GRAY = colors.HexColor("#5A5A5D")
PDF_TEXT = colors.HexColor("#162130")
PDF_MUTED = colors.HexColor("#5E6A75")
VALID_TEMPLATES = {"professional", "modern", "compact"}


def markdown_to_docx(markdown: str, template: str = "professional") -> bytes:
    template_name = normalize_template(template)
    document = Document()
    _configure_docx_document(document)
    blocks = _parse_markdown(markdown)

    _add_docx_brand_header(document, template_name)

    for kind, text in blocks:
        if kind == "title":
            paragraph = document.add_paragraph()
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if template_name == "modern" else WD_ALIGN_PARAGRAPH.LEFT
            )
            run = paragraph.add_run(text)
            run.bold = True
            run.font.size = Pt(18 if template_name == "compact" else 20)
            run.font.color.rgb = BRAND_BLUE if template_name == "modern" else BRAND_GRAY
            continue

        if kind == "section":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(8 if template_name == "compact" else 12)
            paragraph.paragraph_format.space_after = Pt(4)
            run = paragraph.add_run(text.upper())
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = BRAND_BLUE
            if template_name != "compact":
                _add_bottom_border(paragraph)
            continue

        if kind == "subsection":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(4 if template_name == "compact" else 6)
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run(text)
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = BRAND_GRAY
            continue

        if kind == "bullet":
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(0 if template_name == "compact" else 1)
            run = paragraph.add_run(text)
            run.font.size = Pt(10 if template_name == "compact" else 10.5)
            run.font.color.rgb = BRAND_GRAY
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2 if template_name == "compact" else 4)
        run = paragraph.add_run(text)
        run.font.size = Pt(10 if template_name == "compact" else 10.5)
        run.font.color.rgb = BRAND_GRAY

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def markdown_to_pdf(markdown: str, template: str = "professional") -> bytes:
    template_name = normalize_template(template)
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=LETTER)
    width, height = LETTER
    left_margin = 44 if template_name == "compact" else 50
    right_margin = width - 50
    top_margin = height - 54
    bottom_margin = 50
    y = top_margin
    blocks = _parse_markdown(markdown)

    def ensure_space(required_height: float) -> float:
        nonlocal y
        if y - required_height < bottom_margin:
            pdf.showPage()
            y = top_margin
            _draw_pdf_header(pdf, width, height, template_name)
        return y

    _draw_pdf_header(pdf, width, height, template_name)
    y -= 28 if template_name == "compact" else 40

    for kind, text in blocks:
        if kind == "title":
            ensure_space(30)
            pdf.setFillColor(PDF_BRAND_BLUE if template_name == "modern" else PDF_BRAND_GRAY)
            pdf.setFont("Helvetica-Bold", 16 if template_name == "compact" else 18)
            if template_name == "modern":
                pdf.drawCentredString(width / 2, y, text)
            else:
                pdf.drawString(left_margin, y, text)
            y -= 18 if template_name == "compact" else 24
            continue

        if kind == "section":
            ensure_space(24 if template_name == "compact" else 28)
            pdf.setFillColor(PDF_BRAND_BLUE)
            pdf.setFont("Helvetica-Bold", 10)
            if template_name == "modern":
                pdf.setStrokeColor(PDF_BRAND_BLUE)
                pdf.setLineWidth(1)
                pdf.line(left_margin, y + 2, right_margin, y + 2)
                pdf.drawCentredString(width / 2, y - 10, text.upper())
            else:
                if template_name != "compact":
                    pdf.setStrokeColor(PDF_BRAND_BLUE)
                    pdf.setLineWidth(1)
                    pdf.line(left_margin, y + 2, right_margin, y + 2)
                pdf.drawString(left_margin, y - 10, text.upper())
            y -= 16 if template_name == "compact" else 22
            continue

        if kind == "subsection":
            ensure_space(14 if template_name == "compact" else 18)
            pdf.setFillColor(PDF_BRAND_GRAY)
            pdf.setFont("Helvetica-Bold", 11)
            pdf.drawString(left_margin, y, text)
            y -= 12 if template_name == "compact" else 16
            continue

        if kind == "bullet":
            y = _draw_pdf_wrapped_text(
                pdf,
                text,
                left_margin + 12,
                y,
                right_margin - left_margin - 12,
                prefix="* ",
                font_name="Helvetica",
                font_size=10 if template_name == "compact" else 10.5,
                color=PDF_TEXT,
                bottom_margin=bottom_margin,
                top_margin=top_margin,
                width=width,
                height=height,
                template=template_name,
            )
            y -= 1 if template_name == "compact" else 2
            continue

        y = _draw_pdf_wrapped_text(
            pdf,
            text,
            left_margin,
            y,
            right_margin - left_margin,
            prefix="",
            font_name="Helvetica",
            font_size=10 if template_name == "compact" else 10.5,
            color=PDF_TEXT,
            bottom_margin=bottom_margin,
            top_margin=top_margin,
            width=width,
            height=height,
            template=template_name,
        )
        y -= 2 if template_name == "compact" else 4

    pdf.save()
    return buffer.getvalue()


def _configure_docx_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10.5)


def _add_docx_brand_header(document: Document, template: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if template == "modern" else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run("Torilaure")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = BRAND_BLUE

    sub = paragraph.add_run(" E-systems")
    sub.font.size = Pt(16)
    sub.font.color.rgb = BRAND_GRAY

    caption = document.add_paragraph()
    caption.paragraph_format.space_after = Pt(6 if template == "compact" else 10)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER if template == "modern" else WD_ALIGN_PARAGRAPH.LEFT
    caption_text = {
        "professional": "Resume Co-Pilot Export",
        "modern": "Modern Resume Template",
        "compact": "Compact Resume Template",
    }[template]
    cap_run = caption.add_run(caption_text)
    cap_run.italic = True
    cap_run.font.size = Pt(9)
    cap_run.font.color.rgb = BRAND_GRAY


def _add_bottom_border(paragraph) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    paragraph_borders = paragraph_properties.find(qn("w:pBdr"))
    if paragraph_borders is None:
        paragraph_borders = OxmlElement("w:pBdr")
        paragraph_properties.append(paragraph_borders)

    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2A93D5")
    paragraph_borders.append(bottom)


def _draw_pdf_header(pdf: canvas.Canvas, width: float, height: float, template: str) -> None:
    if template == "compact":
        pdf.setFillColor(PDF_BRAND_GRAY)
        pdf.setFont("Helvetica-Bold", 11)
        pdf.drawString(44, height - 30, "Torilaure E-systems")
        pdf.setFillColor(PDF_MUTED)
        pdf.setFont("Helvetica-Oblique", 8)
        pdf.drawRightString(width - 44, height - 30, "Compact Resume Template")
        return

    pdf.setFillColor(PDF_BRAND_BLUE)
    pdf.rect(0, height - 36, width, 36, stroke=0, fill=1)
    pdf.setFillColor(colors.white)
    pdf.setFont("Helvetica-Bold", 16)
    if template == "modern":
        pdf.drawCentredString(width / 2 - 18, height - 23, "Torilaure")
        pdf.setFont("Helvetica", 16)
        pdf.drawCentredString(width / 2 + 44, height - 23, "E-systems")
    else:
        pdf.drawString(50, height - 23, "Torilaure")
        pdf.setFont("Helvetica", 16)
        pdf.drawString(119, height - 23, "E-systems")
    pdf.setFillColor(PDF_MUTED)
    pdf.setFont("Helvetica-Oblique", 8.5)
    pdf.drawRightString(
        width - 50,
        height - 23,
        "Modern Resume Template" if template == "modern" else "Resume Co-Pilot Export",
    )


def _draw_pdf_wrapped_text(
    pdf: canvas.Canvas,
    text: str,
    x: float,
    y: float,
    max_width: float,
    prefix: str,
    font_name: str,
    font_size: float,
    color,
    bottom_margin: float,
    top_margin: float,
    width: float,
    height: float,
    template: str,
) -> float:
    pdf.setFillColor(color)
    indent = stringWidth(prefix, font_name, font_size)
    wrapped_lines = _wrap_pdf_text(text, font_name, font_size, max_width - indent)

    for index, part in enumerate(wrapped_lines):
        if y < bottom_margin:
            pdf.showPage()
            _draw_pdf_header(pdf, width, height, template)
            y = top_margin - (28 if template == "compact" else 40)

        pdf.setFont(font_name, font_size)
        current_prefix = prefix if index == 0 else ""
        pdf.drawString(x, y, f"{current_prefix}{part}")
        y -= font_size + 4

    return y


def _parse_markdown(markdown: str) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []

    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            blocks.append(("title", line[2:].strip()))
        elif line.startswith("## "):
            blocks.append(("section", line[3:].strip()))
        elif line.startswith("### "):
            blocks.append(("subsection", line[4:].strip()))
        elif line.startswith("- "):
            blocks.append(("bullet", line[2:].strip()))
        else:
            blocks.append(("paragraph", line))

    return blocks


def _wrap_pdf_text(text: str, font_name: str, font_size: float, max_width: float) -> list[str]:
    words = text.split()
    if not words:
        return [""]

    lines: list[str] = []
    current = words[0]

    for word in words[1:]:
        trial = f"{current} {word}"
        if stringWidth(trial, font_name, font_size) <= max_width:
            current = trial
        else:
            lines.append(current)
            current = word

    lines.append(current)
    return lines


def normalize_template(template: str) -> str:
    cleaned = template.strip().lower()
    return cleaned if cleaned in VALID_TEMPLATES else "professional"
